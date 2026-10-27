"""Write bilingual (FI/EN) Excel workbook matching example tab structure."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from payloq_schemas.models import ScenarioPack


def write_scenario_xlsx(pack: ScenarioPack, path: Path | str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    _scenario_sheet(wb.active, pack)
    _income_sheet(wb.create_sheet("Income statement"), pack)
    _balance_sheet(wb.create_sheet("Balance sheet"), pack)
    _history_sheet(wb.create_sheet("Balance history"), pack)
    _invoices_sheet(wb.create_sheet("Open invoices"), pack)
    _projection_sheet(wb.create_sheet("Projection"), pack)

    wb.save(path)
    return path


def _scenario_sheet(ws, pack: ScenarioPack) -> None:
    ws.title = "Scenario"
    inp = pack.input
    rows = [
        ("Kenttä / Field", "FI", "EN / Value"),
        ("Kuvaus / Description", inp.scenario_description.fi, inp.scenario_description.en),
        ("Toimiala / Industry", inp.options.industry, inp.options.industry),
        ("Yrityksen koko / Company size", inp.options.company_size, inp.options.company_size),
        ("Kuukaudet / Months", inp.options.months, inp.options.months),
        ("Kieli / Language", inp.options.language, inp.options.language),
        ("Valuutta / Currency", inp.options.currency, inp.options.currency),
        ("Skenaario / Scenario", inp.options.scenario_id or "", inp.meta.scenario or ""),
        (
            "Kuukausiliikevaihto alussa / Monthly revenue start (EUR)",
            inp.parameters_from_language_model.monthly_revenue_start_eur,
            inp.parameters_from_language_model.monthly_revenue_start_eur,
        ),
        (
            "Kasvu / Growth",
            inp.parameters_from_language_model.monthly_revenue_growth,
            inp.parameters_from_language_model.monthly_revenue_growth,
        ),
        (
            "DSO (pv) / Days sales outstanding",
            inp.parameters_from_language_model.days_sales_outstanding,
            inp.parameters_from_language_model.days_sales_outstanding,
        ),
        (
            "DPO (pv) / Days payables outstanding",
            inp.parameters_from_language_model.days_payables_outstanding,
            inp.parameters_from_language_model.days_payables_outstanding,
        ),
        (
            "Kassatavoite / Target cash at end (EUR)",
            inp.parameters_from_language_model.target_cash_at_end_eur,
            inp.parameters_from_language_model.target_cash_at_end_eur,
        ),
        (
            "Erityisinvestointi / Special investment (EUR)",
            inp.parameters_from_language_model.special_investment_eur,
            inp.parameters_from_language_model.special_investment_eur,
        ),
        ("Huomiot / Notes", inp.parameters_from_language_model.notes or "", ""),
    ]
    for r in rows:
        ws.append(list(r))


def _income_sheet(ws, pack: ScenarioPack) -> None:
    headers = [
        "Kuukausi / Month",
        "Liikevaihto / Revenue",
        "Aineet ja palvelut / Materials and services",
        "Henkilöstökulut / Personnel costs",
        "Liiketoiminnan muut kulut / Other operating costs",
        "Poistot / Depreciation",
        "Liikevoitto / Operating profit",
        "Korkokulut / Interest expense",
        "Tulos ennen veroja / Profit before tax",
        "Tulovero / Income tax",
        "Tilikauden tulos / Net result",
    ]
    ws.append(headers)
    for m in pack.monthly_reports.months:
        is_ = m.income_statement
        ws.append(
            [
                m.month,
                is_.revenue,
                is_.materials_and_services,
                is_.personnel_costs,
                is_.other_operating_costs,
                is_.depreciation,
                is_.operating_profit,
                is_.interest_expense,
                is_.profit_before_tax,
                is_.income_tax,
                is_.net_result,
            ]
        )


def _balance_sheet(ws, pack: ScenarioPack) -> None:
    headers = [
        "Kuukausi / Month",
        "Pysyvät vastaavat / Fixed assets",
        "Vaihto-omaisuus / Inventory",
        "Saamiset / Receivables",
        "Rahat / Cash",
        "Vastaavaa yhteensä / Total assets",
        "Oma pääoma / Equity",
        "Lainat / Loans",
        "Ostovelat / Payables",
        "Muut lyhytaikaiset velat / Other current liabilities",
        "Vastattavaa yhteensä / Total equity and liabilities",
    ]
    ws.append(headers)
    for m in pack.monthly_reports.months:
        a = m.balance_sheet.assets
        e = m.balance_sheet.equity_and_liabilities
        ws.append(
            [
                m.month,
                a.fixed_assets,
                a.inventory,
                a.receivables,
                a.cash,
                a.total_assets,
                e.equity,
                e.loans,
                e.payables,
                e.other_current_liabilities,
                e.total_equity_and_liabilities,
            ]
        )


def _history_sheet(ws, pack: ScenarioPack) -> None:
    ws.append(["Päivä / Date", "Saldo / Balance"])
    for p in pack.cash_and_invoices.balance_history:
        ws.append([p.date, p.balance])


def _invoices_sheet(ws, pack: ScenarioPack) -> None:
    ws.append(
        [
            "Tunnus / Id",
            "Tyyppi / Type",
            "Vastapuoli / Counterparty",
            "Kuvaus / Description",
            "Laskutuspäivä / Issue date",
            "Eräpäivä / Due date",
            "Summa sis. ALV / Amount incl. VAT",
            "Tila / Status",
        ]
    )
    for inv in pack.cash_and_invoices.open_invoices:
        ws.append(
            [
                inv.id,
                inv.type,
                inv.counterparty,
                inv.description,
                inv.issue_date,
                inv.due_date,
                inv.amount_incl_vat,
                inv.status,
            ]
        )


def _projection_sheet(ws, pack: ScenarioPack) -> None:
    ws.append(
        [
            "Viikko / Week",
            "Viikon loppu / Week end",
            "Myyntitulot / Sales inflow",
            "Ostoimenot / Purchase outflow",
            "Ennustettu saldo / Projected balance",
        ]
    )
    for w in pack.cash_and_invoices.projection_from_open_invoices:
        ws.append(
            [w.week, w.week_end, w.sales_inflow, w.purchase_outflow, w.projected_balance]
        )
